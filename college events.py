import tkinter as tk
from tkinter import ttk
from openpyxl import Workbook, load_workbook
import os

file_name = "college_events.xlsx"

if not os.path.exists(file_name):

    wb = Workbook()
    ws = wb.active

    ws.append([
        "Event ID",
        "Event Name",
        "Date",
        "Venue",
        "Organizer",
        "Category"
    ])

    wb.save(file_name)

root = tk.Tk()
root.title("College Events Management System")
root.geometry("900x600")

def clear_screen():

    for widget in root.winfo_children():
        widget.destroy()

def login():

    clear_screen()

    tk.Label(
        root,
        text="College Events Management System",
        font=("Arial", 20)
    ).pack(pady=30)

    tk.Label(
        root,
        text="Username"
    ).pack()

    username = tk.Entry(root)
    username.pack(pady=5)

    tk.Label(
        root,
        text="Password"
    ).pack()

    password = tk.Entry(
        root,
        show="*"
    )
    password.pack(pady=5)

    result = tk.Label(root, text="")
    result.pack(pady=10)


    def check_login():

        if username.get() == "admin" and password.get() == "1234":

            dashboard()

        else:

            result.config(
                text="Wrong Username or Password"
            )


    tk.Button(
        root,
        text="Login",
        command=check_login
    ).pack(pady=10)

def dashboard():

    clear_screen()

    tk.Label(
        root,
        text="College Events Management System",
        font=("Arial", 20)
    ).pack(pady=20)

    tk.Label(
        root,
        text="Dashboard",
        font=("Arial", 16)
    ).pack(pady=10)


    tk.Button(
        root,
        text="Add Event",
        width=20,
        command=add_event
    ).pack(pady=5)


    tk.Button(
        root,
        text="View Events",
        width=20,
        command=view_events
    ).pack(pady=5)


    tk.Button(
        root,
        text="Search Event",
        width=20,
        command=search_event
    ).pack(pady=5)


    tk.Button(
        root,
        text="Update Event",
        width=20,
        command=update_event
    ).pack(pady=5)


    tk.Button(
        root,
        text="Delete Event",
        width=20,
        command=delete_event
    ).pack(pady=5)


    tk.Button(
        root,
        text="Logout",
        width=20,
        command=login
    ).pack(pady=15)

def add_event():

    clear_screen()

    tk.Label(
        root,
        text="Add College Event",
        font=("Arial", 20)
    ).pack(pady=20)


    frame = tk.Frame(root)
    frame.pack()


    tk.Label(
        frame,
        text="Event ID"
    ).grid(row=0, column=0, pady=5)

    event_id = tk.Entry(frame)
    event_id.grid(row=0, column=1)


    tk.Label(
        frame,
        text="Event Name"
    ).grid(row=1, column=0, pady=5)

    event_name = tk.Entry(frame)
    event_name.grid(row=1, column=1)


    tk.Label(
        frame,
        text="Date"
    ).grid(row=2, column=0, pady=5)

    date = tk.Entry(frame)
    date.grid(row=2, column=1)


    tk.Label(
        frame,
        text="Venue"
    ).grid(row=3, column=0, pady=5)

    venue = tk.Entry(frame)
    venue.grid(row=3, column=1)


    tk.Label(
        frame,
        text="Organizer"
    ).grid(row=4, column=0, pady=5)

    organizer = tk.Entry(frame)
    organizer.grid(row=4, column=1)


    tk.Label(
        frame,
        text="Category"
    ).grid(row=5, column=0, pady=5)

    category = tk.Entry(frame)
    category.grid(row=5, column=1)


    result = tk.Label(
        root,
        text=""
    )
    result.pack(pady=10)


    def save_event():

        wb = load_workbook(file_name)
        ws = wb.active

        ws.append([
            event_id.get(),
            event_name.get(),
            date.get(),
            venue.get(),
            organizer.get(),
            category.get()
        ])

        wb.save(file_name)

        result.config(
            text="Event Saved Successfully"
        )


    def clear_data():

        event_id.delete(0, tk.END)
        event_name.delete(0, tk.END)
        date.delete(0, tk.END)
        venue.delete(0, tk.END)
        organizer.delete(0, tk.END)
        category.delete(0, tk.END)


    tk.Button(
        root,
        text="Save",
        command=save_event
    ).pack(pady=5)


    tk.Button(
        root,
        text="Clear",
        command=clear_data
    ).pack(pady=5)


    tk.Button(
        root,
        text="Back",
        command=dashboard
    ).pack(pady=10)

def view_events():

    clear_screen()

    tk.Label(
        root,
        text="College Events",
        font=("Arial", 20)
    ).pack(pady=20)


    columns = (
        "ID",
        "Event Name",
        "Date",
        "Venue",
        "Organizer",
        "Category"
    )


    table = ttk.Treeview(
        root,
        columns=columns,
        show="headings"
    )


    table.heading(
        "ID",
        text="Event ID"
    )

    table.heading(
        "Event Name",
        text="Event Name"
    )

    table.heading(
        "Date",
        text="Date"
    )

    table.heading(
        "Venue",
        text="Venue"
    )

    table.heading(
        "Organizer",
        text="Organizer"
    )

    table.heading(
        "Category",
        text="Category"
    )


    table.column("ID", width=80)
    table.column("Event Name", width=150)
    table.column("Date", width=100)
    table.column("Venue", width=150)
    table.column("Organizer", width=150)
    table.column("Category", width=120)


    table.pack(
        fill="both",
        expand=True,
        padx=10,
        pady=10
    )


    wb = load_workbook(file_name)
    ws = wb.active


    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        table.insert(
            "",
            "end",
            values=row
        )


    tk.Button(
        root,
        text="Back",
        command=dashboard
    ).pack(pady=10)


# ---------------- SEARCH EVENT ----------------

def search_event():

    clear_screen()

    tk.Label(
        root,
        text="Search Event",
        font=("Arial", 20)
    ).pack(pady=20)


    tk.Label(
        root,
        text="Enter Event ID or Event Name"
    ).pack()


    search = tk.Entry(root)
    search.pack(pady=5)


    result = tk.Label(
        root,
        text=""
    )
    result.pack(pady=5)


    columns = (
        "ID",
        "Event Name",
        "Date",
        "Venue",
        "Organizer",
        "Category"
    )


    table = ttk.Treeview(
        root,
        columns=columns,
        show="headings"
    )


    for column in columns:

        table.heading(
            column,
            text=column
        )

        table.column(
            column,
            width=130
        )


    table.pack(
        fill="both",
        expand=True,
        pady=10
    )


    def search_data():

        # Clear previous result

        for item in table.get_children():

            table.delete(item)


        search_value = search.get().lower()


        wb = load_workbook(file_name)
        ws = wb.active


        found = False


        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            event_id = str(row[0]).lower()
            event_name = str(row[1]).lower()


            if (
                search_value == event_id
                or
                search_value == event_name
            ):

                table.insert(
                    "",
                    "end",
                    values=row
                )

                found = True


        if found:

            result.config(
                text="Event Found"
            )

        else:

            result.config(
                text="Event Not Found"
            )


    tk.Button(
        root,
        text="Search",
        command=search_data
    ).pack(pady=5)


    tk.Button(
        root,
        text="Back",
        command=dashboard
    ).pack(pady=5)

def update_event():

    clear_screen()

    tk.Label(
        root,
        text="Update Event",
        font=("Arial", 20)
    ).pack(pady=20)


    frame = tk.Frame(root)
    frame.pack()


    tk.Label(
        frame,
        text="Event ID"
    ).grid(row=0, column=0, pady=5)

    event_id = tk.Entry(frame)
    event_id.grid(row=0, column=1)


    tk.Label(
        frame,
        text="Event Name"
    ).grid(row=1, column=0, pady=5)

    event_name = tk.Entry(frame)
    event_name.grid(row=1, column=1)


    tk.Label(
        frame,
        text="Date"
    ).grid(row=2, column=0, pady=5)

    date = tk.Entry(frame)
    date.grid(row=2, column=1)


    tk.Label(
        frame,
        text="Venue"
    ).grid(row=3, column=0, pady=5)

    venue = tk.Entry(frame)
    venue.grid(row=3, column=1)


    tk.Label(
        frame,
        text="Organizer"
    ).grid(row=4, column=0, pady=5)

    organizer = tk.Entry(frame)
    organizer.grid(row=4, column=1)


    tk.Label(
        frame,
        text="Category"
    ).grid(row=5, column=0, pady=5)

    category = tk.Entry(frame)
    category.grid(row=5, column=1)


    result = tk.Label(
        root,
        text=""
    )
    result.pack(pady=10)


    def find_event():

        wb = load_workbook(file_name)
        ws = wb.active


        for row in ws.iter_rows(
            min_row=2
        ):

            if str(row[0].value).lower() == event_id.get().lower():

                event_name.delete(0, tk.END)
                event_name.insert(0, row[1].value)

                date.delete(0, tk.END)
                date.insert(0, row[2].value)

                venue.delete(0, tk.END)
                venue.insert(0, row[3].value)

                organizer.delete(0, tk.END)
                organizer.insert(0, row[4].value)

                category.delete(0, tk.END)
                category.insert(0, row[5].value)

                result.config(
                    text="Event Found"
                )

                return


        result.config(
            text="Event Not Found"
        )


    def update_data():

        wb = load_workbook(file_name)
        ws = wb.active


        for row in ws.iter_rows(
            min_row=2
        ):

            if str(row[0].value).lower() == event_id.get().lower():

                row[1].value = event_name.get()
                row[2].value = date.get()
                row[3].value = venue.get()
                row[4].value = organizer.get()
                row[5].value = category.get()


                wb.save(file_name)


                result.config(
                    text="Event Updated Successfully"
                )

                return


        result.config(
            text="Event Not Found"
        )


    tk.Button(
        root,
        text="Find Event",
        command=find_event
    ).pack(pady=5)


    tk.Button(
        root,
        text="Update",
        command=update_data
    ).pack(pady=5)


    tk.Button(
        root,
        text="Back",
        command=dashboard
    ).pack(pady=10)

def delete_event():

    clear_screen()

    tk.Label(
        root,
        text="Delete Event",
        font=("Arial", 20)
    ).pack(pady=20)


    tk.Label(
        root,
        text="Enter Event ID"
    ).pack()


    event_id = tk.Entry(root)
    event_id.pack(pady=5)


    result = tk.Label(
        root,
        text=""
    )
    result.pack(pady=10)


    def delete_data():

        wb = load_workbook(file_name)
        ws = wb.active


        for row in range(
            2,
            ws.max_row + 1
        ):

            if str(ws.cell(row, 1).value).lower() == event_id.get().lower():

                ws.delete_rows(
                    row,
                    1
                )

                wb.save(file_name)


                result.config(
                    text="Event Deleted Successfully"
                )

                return


        result.config(
            text="Event Not Found"
        )


    tk.Button(
        root,
        text="Delete",
        command=delete_data
    ).pack(pady=5)


    tk.Button(
        root,
        text="Back",
        command=dashboard
    ).pack(pady=10)

login()

root.mainloop()